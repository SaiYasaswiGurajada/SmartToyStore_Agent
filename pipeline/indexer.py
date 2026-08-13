"""
pipeline/indexer.py — Parses uploaded files and the knowledge base into chunks.

Chunking strategy: subsection level (### N.N) for the KB, paragraph-level for
other documents. KB sections 8 and 9 are NEVER indexed (they are behaviour specs).
"""

from __future__ import annotations
from typing import Optional
import re
import io
from pathlib import Path
from dataclasses import dataclass, field

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from config.config import (
    CHUNK_ID_PATTERN, INDEXABLE_KB_SECTIONS, PLACEHOLDER_PATTERN,
    ALLOWED_UPLOAD_EXTENSIONS, MAX_UPLOAD_MB, KB_DIR,
)

# --------------------------------------------------------------------------
# Data structure
# --------------------------------------------------------------------------

@dataclass
class Chunk:
    text: str
    source: str           # filename or KB path
    subsection: str = "" # "1.1", "2.3", etc. — empty for non-KB docs
    placeholder_flagged: bool = False
    embedding: list[float] = field(default_factory=list)


# --------------------------------------------------------------------------
# Placeholder detection
# --------------------------------------------------------------------------

_PLACEHOLDER_RE = re.compile(PLACEHOLDER_PATTERN)
_SUBSECTION_RE = re.compile(CHUNK_ID_PATTERN, re.MULTILINE)
_SECTION_HEADER_RE = re.compile(r"^##\s+(\d+)\.", re.MULTILINE)


def _flag_placeholders(chunks: list[Chunk]) -> list[Chunk]:
    for chunk in chunks:
        if _PLACEHOLDER_RE.search(chunk.text):
            chunk.placeholder_flagged = True
    return chunks


# --------------------------------------------------------------------------
# KB chunker — subsection level
# --------------------------------------------------------------------------

def _get_section_number(text_before: str) -> Optional[int]:
    """Return the section number of the most recent ## N. header."""
    matches = list(_SECTION_HEADER_RE.finditer(text_before))
    if not matches:
        return None
    return int(matches[-1].group(1))


def chunk_knowledge_base(kb_path: Path) -> list[Chunk]:
    """
    Split the KB file at ### N.N subsection boundaries.
    Only index sections 1–7; skip 8 and 9.
    """
    text = kb_path.read_text(encoding="utf-8")
    chunks: list[Chunk] = []

    # Split on subsection headings (### N.N ...)
    # We keep the heading inside the chunk text
    parts = re.split(r"(?=^###\s+\d+\.\d+)", text, flags=re.MULTILINE)

    for part in parts:
        part = part.strip()
        if not part:
            continue

        # Extract subsection id
        m = _SUBSECTION_RE.match(part)
        if not m:
            continue  # preamble / section headers without ###

        subsection_id = m.group(1)  # e.g. "1.1"
        section_num = int(subsection_id.split(".")[0])

        if section_num not in INDEXABLE_KB_SECTIONS:
            continue  # skip sections 8 and 9

        chunk = Chunk(
            text=part,
            source=str(kb_path.name),
            subsection=subsection_id,
        )
        chunks.append(chunk)

    return _flag_placeholders(chunks)


# --------------------------------------------------------------------------
# Generic document chunkers
# --------------------------------------------------------------------------

def _chunk_text_by_paragraph(text: str, source: str) -> list[Chunk]:
    """Split plain text on double-newlines, ~300 words max per chunk."""
    paragraphs = re.split(r"\n\s*\n", text.strip())
    chunks: list[Chunk] = []
    current = []
    current_words = 0

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        words = len(para.split())
        if current_words + words > 300 and current:
            chunks.append(Chunk(text="\n\n".join(current), source=source))
            current = []
            current_words = 0
        current.append(para)
        current_words += words

    if current:
        chunks.append(Chunk(text="\n\n".join(current), source=source))

    return _flag_placeholders(chunks)


def parse_txt(path: Path) -> list[Chunk]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return _chunk_text_by_paragraph(text, path.name)


def parse_md(path: Path) -> list[Chunk]:
    # If it's the main KB, use subsection chunker
    if path.parent == KB_DIR and path.name == "smart_toy_store_knowledge_base.md":
        return chunk_knowledge_base(path)
    text = path.read_text(encoding="utf-8", errors="replace")
    return _chunk_text_by_paragraph(text, path.name)


def parse_pdf(path: Path) -> list[Chunk]:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        pages_text = [page.get_text() for page in doc]
        full_text = "\n\n".join(pages_text)
        return _chunk_text_by_paragraph(full_text, path.name)
    except ImportError:
        raise RuntimeError("PyMuPDF (fitz) not installed. Run: pip install pymupdf")


def parse_docx(path: Path) -> list[Chunk]:
    try:
        from docx import Document
        doc = Document(str(path))
        full_text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return _chunk_text_by_paragraph(full_text, path.name)
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")


def parse_file(path: Path) -> list[Chunk]:
    """Dispatch to the right parser based on extension."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    elif suffix == ".docx":
        return parse_docx(path)
    elif suffix in (".txt",):
        return parse_txt(path)
    elif suffix in (".md",):
        return parse_md(path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


# --------------------------------------------------------------------------
# Ingest the default knowledge base files at startup
# --------------------------------------------------------------------------

def load_default_kb() -> list[Chunk]:
    """Load smart_toy_store_knowledge_base.md and FAQ.txt from kb/."""
    chunks: list[Chunk] = []
    kb_md = KB_DIR / "smart_toy_store_knowledge_base.md"
    faq_txt = KB_DIR / "FAQ.txt"

    if kb_md.exists():
        chunks.extend(chunk_knowledge_base(kb_md))
    if faq_txt.exists():
        chunks.extend(parse_txt(faq_txt))

    return chunks
